import json, pandas as pd
from decimal import Decimal, ROUND_HALF_UP
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def r(v, n=3): return str(Decimal(repr(round(float(v), 8))).quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP))
m = pd.read_csv('outputs/metrics.csv').set_index('Model')
meta = json.load(open('models/meta.json')); ex = json.load(open('models/holdout_example.json'))
best = meta['model']; lg, xg, rf, cb, bl = (m.loc[k] for k in ['LightGBM', 'XGBoost', 'Random Forest', 'CatBoost', 'Baseline: item spike rate'])

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(11)
H = lambda t, l=1: doc.add_heading(t, level=l)
P = lambda t: doc.add_paragraph(t)
B = lambda t: doc.add_paragraph(t, style='List Bullet')
def IMG(p, w=6.2, cap=None):
    doc.add_picture(p, width=Inches(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        c = P(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.runs[0].italic = True; c.runs[0].font.size = Pt(9)
def TABLE(header, rows):
    t = doc.add_table(rows=1, cols=len(header)); t.style = 'Light Grid Accent 1'
    for i, h in enumerate(header): t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)

doc.add_heading('Early Warning of Food Price Spikes in Sri Lankan Markets', 0)
P('Hands-on Activity: ML Model Deployment & Comparison (Random Forest, XGBoost, LightGBM, CatBoost, with SHAP analysis)')
P('Name: [YOUR NAME]     Index number: [YOUR INDEX NUMBER]     Course: [COURSE CODE]     Date: [DATE]')

H('1. Problem definition')
P('Food is a large part of household spending in Sri Lanka, and prices can change sharply from one month to the next. Sudden rises hurt low-income families most, '
  'and the 2022 economic crisis showed how quickly food prices can climb. If a rise can be anticipated, households, traders and the authorities have time to respond '
  '(for example by releasing stocks, adjusting imports or targeting support).')
P('This activity builds an early-warning model that answers one question: will the retail price of a food item, in a given Sri Lankan market, rise by 10% or more next month? '
  'This is a binary classification task (spike = yes or no). Four ensemble models are trained, compared and the best one is selected. '
  'SHAP is then used to explain what drives the predictions, and the selected model is deployed behind a FastAPI backend and a Streamlit dashboard.')

H('2. Dataset')
P('Source: retail food prices collected by the UN World Food Programme (WFP) and published on the Humanitarian Data Exchange '
  '(https://data.humdata.org/dataset/wfp-food-prices-for-sri-lanka). The file used has 32,496 price records for 44 markets and 42 items, in Sri Lankan rupees (LKR), '
  'from January 2004 to September 2025. Items include rice, wheat flour, sugar, vegetables, fruit, fish, eggs, pulses and fuel.')
P('Coverage is thin before mid-2023 (a few hundred records per year in most years) and dense from June 2023 (about 1,000 records every month), as Figure 1 shows.')
IMG('outputs/eda_records_per_month.png', 6.2, 'Figure 1: Price records per month in the WFP file')
H('2.1 Cleaning', 2)
B('Kept retail prices of food items in real markets. Removed 194 "National Average" rows (no location), wholesale prices and fuel (non-food).')
B('Removed items with fewer than 100 records, and kept one record per market, item and month. This leaves 31,246 records, 37 items and 42 markets.')
B('Kept only records with a full 3-month price history for the same market and item (needed for the features), giving 22,575 rows. About 19% of them are spikes.')
IMG('outputs/eda_spike_rate_by_item.png', 4.6, 'Figure 2: How often each item rises 10% or more in a month')
P('Figure 2 shows that vegetables such as tomatoes, carrots and eggplant spike often, while staples such as rice, sugar and wheat flour rarely do. '
  'Because "which item is it" is such a strong clue, a baseline that simply uses each item\'s usual spike rate is included to check what the models add.')

H('3. Methodology')
H('3.1 Features (past information only)', 2)
P('All features use prices up to the latest known month, never the month being predicted:')
B('price_now: latest price (LKR); change_1m and change_2m: price change over the last two months; vs_3m_avg: latest price compared with its 3-month average.')
B('national_avg_now and national_change_1m: the latest national average for the item and how it moved; market_vs_national: whether this market is dearer or cheaper than the national average.')
B('month (season), latitude, longitude, province (9 columns) and item (37 columns, one per item). In total 56 features.')
P('Target: spike = 1 if next month\'s price is at least 10% higher than the latest price.')
H('3.2 Splitting by time', 2)
B('Train: target months up to June 2024 (11,163 rows, 19.4% spikes). Models learn and are tuned here.')
B('Validation: July to December 2024 (4,527 rows, 20.8% spikes). Used to choose the best model and the alert threshold.')
B('Test: January to September 2025 (6,884 rows, 17.1% spikes). Used once for the final scores.')
B(f'One row was held out for the inference demo (Tomatoes, Colombo City, September 2025) and removed from all three sets.')
P('The split is by time, not random, because a real early-warning system only sees the past. A random split would let the model learn from months after the ones it is tested on.')
H('3.3 Models and tuning', 2)
B('Random Forest: many trees trained independently, predictions averaged.')
B('XGBoost: trees built one after another, each correcting the errors of the previous ones.')
B('LightGBM: a faster boosting library that grows trees leaf by leaf.')
B('CatBoost: a boosting library with built-in protection against over-fitting.')
B('Each model was tuned with RandomizedSearchCV and 3-fold TimeSeriesSplit (each fold validates on a period after its training data), scored by ROC-AUC, on the training period only.')
B('Baseline: each item\'s spike rate in the training period.')
H('3.4 Evaluation metrics', 2)
B('ROC-AUC: the chance that a real spike is ranked above a non-spike (0.5 = guessing). Used to choose the best model, because it does not depend on a threshold.')
B('PR-AUC (average precision): how well spikes are found while keeping false alarms low. Guessing gives about the spike rate (0.17 on the test set).')
B('Precision, recall and F1 at an alert threshold chosen on the validation set. Plain accuracy is not used: always predicting "no spike" would be about 83% accurate and useless.')

H('4. Results')
rows = []
for name in ['LightGBM', 'XGBoost', 'Random Forest', 'CatBoost', 'Baseline: item spike rate']:
    x = m.loc[name]
    rows.append([name, r(x.Val_AUC), r(x.Test_AUC), r(x.Test_PR_AUC), r(x.Test_Precision), r(x.Test_Recall), r(x.Test_F1), '-' if x.Train_seconds == 0 else r(x.Train_seconds, 1)])
TABLE(['Model', 'Val AUC', 'Test AUC', 'Test PR-AUC', 'Precision', 'Recall', 'F1', 'Tuning time (s)'], rows)
P('')
IMG('outputs/model_comparison.png', 5.8, 'Figure 3: ROC-AUC on validation (used to choose) and test (final score)')
IMG('outputs/roc_pr_curves.png', 6.3, 'Figure 4: ROC and precision-recall curves on the test period')
H('4.1 Model selection', 2)
P(f'LightGBM had the highest validation ROC-AUC ({r(lg.Val_AUC)}), with XGBoost very close ({r(xg.Val_AUC)}), so LightGBM was selected. On the test period it scored ROC-AUC {r(lg.Test_AUC)}, '
  f'PR-AUC {r(lg.Test_PR_AUC)} (the spike rate is {r(meta["spike_rate_test"])}), precision {r(lg.Test_Precision)} and recall {r(lg.Test_Recall)}.')
IMG('outputs/confusion_matrix.png', 3.4, 'Figure 5: LightGBM confusion matrix on the test period')
P('At the chosen threshold the model raised 2,612 alerts in the test period. 716 of them were real spikes (precision 0.27), and it caught 716 of the 1,178 real spikes (recall 0.61). '
  'In other words it catches about 6 in 10 spikes, at the price of roughly 3 false alarms for every real spike found.')
H('4.2 Is the best model really better?', 2)
P('A bootstrap over whole test months (1,000 resamples) was used to see whether the ROC-AUC gaps could be chance:')
TABLE(['Comparison (A vs B)', 'AUC(A) - AUC(B)', '95% range', 'Clear difference?'], [
    ['LightGBM vs Baseline', '0.031', '-0.038 to 0.109', 'no'], ['LightGBM vs Random Forest', '0.021', '-0.008 to 0.045', 'no'],
    ['LightGBM vs XGBoost', '0.002', '-0.012 to 0.016', 'no'], ['LightGBM vs CatBoost', '0.016', '-0.008 to 0.043', 'no']])
P('')
B(f'The four models are very close. LightGBM and XGBoost are practically tied (test AUC {r(lg.Test_AUC)} vs {r(xg.Test_AUC)}), and no gap between the models has a 95% range that excludes zero. '
  'I therefore do not claim LightGBM is truly better than the others; it simply had the best validation score, which was the rule set in advance.')
B(f'All four models beat the baseline on test ROC-AUC ({r(bl.Test_AUC)} for the baseline) and PR-AUC, but only by a small margin, and the bootstrap range for LightGBM vs baseline includes zero. '
  'Much of the predictable part of a spike is simply "which item is it", which the baseline already captures. The models add some skill beyond that, but modestly.')
B('Spikes are partly random (weather, supply shocks, reporting noise), so a perfect early-warning model is not realistic. The value of the model lies more in a ranked risk score and in the explanations below than in exact predictions.')

H('5. SHAP analysis')
P('SHAP (SHapley Additive exPlanations) explains a prediction by giving each feature a value: how much it pushed the spike risk up or down compared with the average prediction. '
  'It was computed for the selected LightGBM model on 2,500 test rows. For LightGBM the values are in log-odds (positive = higher spike risk). Item and province are stored as many 0/1 columns, so they are also added back together.')
IMG('outputs/shap_importance.png', 5.2, 'Figure 6: Average impact of each feature on spike risk')
P('The largest overall influence comes from the national price trend (national_change_1m, about 20% of the total impact), then which item it is (about 14%), '
  'how the market compares with the national price (about 13%), the national price level (about 12%), the price change in the month before last (about 10%) and the calendar month (about 8%).')
IMG('outputs/shap_beeswarm.png', 5.6, 'Figure 7: SHAP summary (each dot is one test row; red = high feature value)')
IMG('outputs/shap_dependence.png', 6.4, 'Figure 8: Shape of the effect of three features')
IMG('outputs/shap_items.png', 4.8, 'Figure 9: Items that raise or lower spike risk')
H('5.1 What the SHAP results say', 2)
B('Mean reversion: when an item\'s price has just fallen (negative change_1m or change_2m), the risk of a spike next month goes up; after a rise, the risk goes down. Prices that dropped sharply tend to bounce back.')
B('Catch-up effect: a market that is cheaper than the national average (negative market_vs_national) has a higher spike risk; a dearer market has a lower risk.')
B('National momentum: when the national average price rose last month, spike risk is higher, so price pressure spreads across markets.')
B('Item type: imported onions, beans, carrots and skipjack tuna carry the highest risk; lentils, white rice, red chili and local potatoes carry the lowest. This matches Figure 2.')
B('Season: January shows the highest risk. The test period covers only January to September 2025, so effects of October to December cannot be judged from the test set.')
B('Cheap items (low national price) spike more than expensive ones, which reflects how volatile cheap vegetables are.')
H('5.2 Explaining one prediction', 2)
IMG('outputs/shap_waterfall.png', 5.4, 'Figure 10: Waterfall for the held-out example (Tomatoes, Colombo City, September 2025)')
P(f'For the held-out example, the latest price was {ex["price_now"]:,.2f} LKR after {ex["price_prev"]:,.2f} and {ex["price_3m"]:,.2f} LKR in the two months before. '
  'The deployed model gave a spike probability of about 45%, so it raised an alert. The main drivers were the item (tomatoes are volatile), the national price rise, and the large fall in the '
  'price two months earlier (mean reversion). The month (September) pushed the risk down. '
  f'The actual price rose only to {ex["actual_next_price"]:,.2f} LKR (+3.3%), so this alert was a false alarm. That is an honest illustration of the precision figure above.')

H('6. System architecture and inference')
IMG('outputs/architecture.png', 6.2, 'Figure 11: Application architecture')
B('Model: LightGBM was retrained on all rows (excluding the held-out row) with its tuned settings and saved with joblib.')
B('Backend (FastAPI): backend/api.py exposes GET /health, GET /model-info, GET /markets, GET /commodities and POST /predict. Input is validated with Pydantic and a SHAP explanation is returned with each prediction.')
B('Frontend (Streamlit): frontend/app.py lets the user choose an item, a market and a month, enter the last three prices and the national averages, and see the spike probability, the alert, '
  'and a SHAP bar chart explaining the result. It also has tabs for the model comparison, the SHAP charts and a price explorer.')
IMG('outputs/dashboard_predict.png', 5.8, 'Figure 12: Dashboard result for the held-out example (served by the FastAPI backend)')

H('7. Limitations and future work')
B('Only past prices are used. Weather, fuel and exchange-rate data, import data and festival calendars could improve the predictions.')
B('The gains over the item-rate baseline are modest and not statistically clear, and there is a single test period of nine months.')
B('Coverage before mid-2023 is thin, so most of the training signal comes from 2023 to 2025.')
B('A 10% threshold is a choice. A different threshold, or predicting the size of the price change (regression), may suit other uses.')
B('The file used is a copy of the WFP data with records up to September 2025. I did not independently verify it against the Humanitarian Data Exchange version, and I did not check how the WFP prices were collected.')
B('The alert threshold favours catching spikes (recall 0.61) over avoiding false alarms (precision 0.27). A different balance can be chosen depending on the cost of each kind of error.')

H('8. Conclusion')
P(f'Four tree-ensemble models were compared for early warning of food price spikes in Sri Lankan markets. LightGBM was selected on validation ROC-AUC ({r(lg.Val_AUC)}) and reached ROC-AUC {r(lg.Test_AUC)} on unseen 2025 data, '
  'only slightly above an item-based baseline and practically tied with XGBoost. SHAP showed that the model relies on sensible signals: the item, recent price falls (mean reversion), '
  'the national price trend and how the market compares with the national price. The model was deployed with a FastAPI backend and a Streamlit dashboard that explains every prediction.')

H('9. Use of AI tools')
P('AI tools (Claude) were used to help plan and write the code and the report draft. I reviewed and ran the code, and I understand the data preparation, the models, the evaluation, SHAP and the application. '
  '[Edit this paragraph so that it describes how you used the tools yourself.]')

H('References')
B('World Food Programme. Sri Lanka - Food Prices. Humanitarian Data Exchange. https://data.humdata.org/dataset/wfp-food-prices-for-sri-lanka')
B('Lundberg, S. and Lee, S.-I. (2017). A unified approach to interpreting model predictions. NeurIPS 2017.')
B('Breiman, L. (2001). Random Forests. Machine Learning, 45, 5-32.')
B('Chen, T. and Guestrin, C. (2016). XGBoost: A scalable tree boosting system. KDD 2016.')
B('Ke, G. et al. (2017). LightGBM: A highly efficient gradient boosting decision tree. NeurIPS 2017.')
B('Prokhorenkova, L. et al. (2018). CatBoost: unbiased boosting with categorical features. NeurIPS 2018.')
doc.save('report/Food_Price_Spike_Report.docx'); print('saved')
